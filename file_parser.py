"""Extract text from candidate submissions (code, documents, notebooks, zip projects) + static checks."""
import ast
import io
import json
import re
import zipfile
from pathlib import Path

import config

TEXT_EXT = {".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rb", ".php", ".sql",
            ".html", ".css", ".json", ".md", ".txt", ".csv", ".yml", ".yaml", ".xml", ".sh", ".r", ".ini", ".toml",
            ".ipynb"}
DOC_EXT = {".pdf", ".docx"}
UPLOAD_TYPES = sorted(e.lstrip(".") for e in (TEXT_EXT | DOC_EXT | {".zip"}))

MAX_FILE_CHARS = 15000
MAX_TOTAL_CHARS = 45000
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
_SKIP_DIRS = ("__pycache__", "node_modules", ".git/", ".venv", "venv/", ".idea", ".vscode")


def _decode(data):
    return data.decode("utf-8", errors="replace")


def read_one(name, data):
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages[:30])
    if ext == ".docx":
        from docx import Document
        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    if ext == ".ipynb":
        try:
            nb = json.loads(_decode(data))
            out = []
            for cell in nb.get("cells", []):
                src = "".join(cell.get("source", []))
                out.append(("# [code cell]\n" if cell.get("cell_type") == "code" else "# [markdown cell]\n") + src)
            return "\n\n".join(out)
        except json.JSONDecodeError:
            return _decode(data)
    return _decode(data)


def expand(name, data):
    """Return [(filename, text)] for one uploaded file (zip archives are expanded safely in memory)."""
    ext = Path(name).suffix.lower()
    if ext != ".zip":
        if ext not in TEXT_EXT | DOC_EXT:
            return []
        return [(name, read_one(name, data))]
    out, budget = [], 20 * 1024 * 1024
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for info in zf.infolist()[:200]:
                n = info.filename
                if info.is_dir() or ".." in n or n.startswith("/") or any(s in n for s in _SKIP_DIRS):
                    continue
                if Path(n).suffix.lower() not in TEXT_EXT | DOC_EXT:
                    continue
                if info.file_size > 2 * 1024 * 1024 or info.file_size > budget:
                    continue
                budget -= info.file_size
                try:
                    out.append((n, read_one(n, zf.read(info))))
                except Exception:  # noqa: BLE001  (a single unreadable file must not break the submission)
                    continue
                if len(out) >= 40:
                    break
    except zipfile.BadZipFile:
        return []
    return out


def analyze_python(name, text):
    """Lightweight static analysis used as evidence for the AI evaluator."""
    findings = []
    try:
        tree = ast.parse(text)
    except SyntaxError as e:
        return [f"{name}: SYNTAX ERROR at line {e.lineno}: {e.msg}"]
    for node in ast.walk(tree):
        if isinstance(node, ast.ExceptHandler) and node.type is None:
            findings.append(f"{name}: bare 'except:' at line {node.lineno} hides errors")
        if isinstance(node, ast.Call) and getattr(node.func, "id", "") in ("eval", "exec"):
            findings.append(f"{name}: use of {node.func.id}() at line {node.lineno} is unsafe")
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            body = [b for b in node.body if not (isinstance(b, ast.Expr) and isinstance(getattr(b, "value", None), ast.Constant))]
            if all(isinstance(b, ast.Pass) for b in body) and body:
                findings.append(f"{name}: function '{node.name}' is an empty stub (line {node.lineno})")
    todos = len(re.findall(r"\b(TODO|FIXME)\b", text))
    if todos:
        findings.append(f"{name}: {todos} TODO/FIXME marker(s) left in code")
    if re.search(r"(api[_-]?key|password|secret|token)\s*=\s*['\"][^'\"]{4,}['\"]", text, re.I):
        findings.append(f"{name}: possible hard-coded secret")
    return findings


def build_bundle(files, notes=""):
    """files: [(filename, bytes)] -> dict(text, file_names, findings, truncated)"""
    sections, names, findings, total, truncated = [], [], [], 0, False
    for fname, data in files:
        try:
            parts = expand(fname, data)
        except Exception as e:  # noqa: BLE001
            findings.append(f"{fname}: could not be read ({e})")
            continue
        if not parts:
            findings.append(f"{fname}: unsupported or empty file, ignored")
        for pname, text in parts:
            names.append(pname)
            if pname.lower().endswith(".py"):
                findings.extend(analyze_python(pname, text))
            if len(text) > MAX_FILE_CHARS:
                text, truncated = text[:MAX_FILE_CHARS] + "\n...[truncated]", True
            if total + len(text) > MAX_TOTAL_CHARS:
                truncated = True
                continue
            total += len(text)
            sections.append(f"=== FILE: {pname} ===\n{text}")
    if notes and notes.strip():
        sections.append("=== CANDIDATE NOTES ===\n" + notes.strip()[:5000])
    return {"text": "\n\n".join(sections), "file_names": names, "findings": findings, "truncated": truncated}


def save_uploads(cid, review_no, files):
    folder = config.UPLOAD_DIR / cid / f"review_{review_no}"
    folder.mkdir(parents=True, exist_ok=True)
    saved = []
    for fname, data in files:
        safe = re.sub(r"[^A-Za-z0-9._-]", "_", Path(fname).name) or "file"
        (folder / safe).write_bytes(data)
        saved.append(safe)
    return saved
