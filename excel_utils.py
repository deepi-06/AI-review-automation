"""Sample Excel generation and candidate import (OpenPyXL + Pandas)."""
import io
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

COLUMNS = ["ID", "Name", "Email", "Role", "Department", "Skills", "Experience"]

SAMPLE_ROWS = [
    ["C001", "Aarav Sharma", "aarav.sharma@example.com", "Python Developer Intern", "Engineering",
     "Python, Flask, SQL, REST APIs, Git", "6 months - built a student portal with Flask"],
    ["C002", "Priya Nair", "priya.nair@example.com", "Data Analyst Intern", "Analytics",
     "Python, Pandas, Excel, Power BI, SQL", "1 year - college analytics club projects"],
    ["C003", "Rohan Mehta", "rohan.mehta@example.com", "Machine Learning Intern", "AI Research",
     "Python, scikit-learn, NumPy, Matplotlib, basic TensorFlow", "8 months - Kaggle competitions and a capstone"],
    ["C004", "Sneha Reddy", "sneha.reddy@example.com", "Frontend Developer Intern", "Engineering",
     "HTML, CSS, JavaScript, React, Figma", "4 months - freelance landing pages"],
    ["C005", "Karthik Subramanian", "karthik.s@example.com", "Java Backend Intern", "Engineering",
     "Java, Spring Boot, MySQL, JUnit, Maven", "1.5 years - academic projects and one internship"],
    ["C006", "Ananya Iyer", "ananya.iyer@example.com", "DevOps Intern", "Infrastructure",
     "Linux, Docker, Git, GitHub Actions, AWS basics", "6 months - home lab and CI projects"],
    ["C007", "Vikram Singh", "vikram.singh@example.com", "QA Automation Intern", "Quality Engineering",
     "Python, Selenium, PyTest, Postman, manual testing", "1 year - freelance testing"],
    ["C008", "Meera Joshi", "meera.joshi@example.com", "Cybersecurity Intern", "Security",
     "Networking, Kali Linux, Wireshark, OWASP Top 10, Python scripting", "5 months - CTF competitions"],
    ["C009", "Arjun Das", "arjun.das@example.com", "Data Engineering Intern", "Data Platform",
     "Python, SQL, PostgreSQL, Airflow basics, ETL", "2 years - academic and part-time projects"],
    ["C010", "Divya Kapoor", "divya.kapoor@example.com", "Full Stack Developer Intern", "Engineering",
     "JavaScript, Node.js, Express, MongoDB, React", "10 months - hackathon projects and open source"],
]


def build_sample_workbook():
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"
    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", start_color="1F4E79")
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font, c.fill = head_font, head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in SAMPLE_ROWS:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(vertical="top", wrap_text=True)
    for col, w in zip("ABCDEFG", [8, 22, 30, 28, 20, 48, 46]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Instructions")
    notes = [
        "How to use this file",
        "1. Keep the 7 column headers on the 'Candidates' sheet exactly as they are.",
        "2. ID must be unique per candidate (e.g. C011). Candidates log in with Email + ID.",
        "3. Replace the example.com addresses with real inboxes if you want to receive demo emails.",
        "4. Upload the file in Admin > Candidates. Re-uploading an existing ID updates that candidate.",
        "Note: the 10 rows are fictional sample data created for demonstration.",
    ]
    for i, t in enumerate(notes, 1):
        info.cell(row=i, column=1, value=t).font = Font(name="Arial", bold=(i == 1), size=11 if i == 1 else 10)
    info.column_dimensions["A"].width = 100
    return wb


def sample_bytes():
    buf = io.BytesIO()
    build_sample_workbook().save(buf)
    return buf.getvalue()


def write_sample(path):
    build_sample_workbook().save(str(path))
    return path


_EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def parse_candidates(file_or_bytes):
    """Return (valid_rows, errors) from an uploaded Excel file."""
    src = io.BytesIO(file_or_bytes) if isinstance(file_or_bytes, (bytes, bytearray)) else file_or_bytes
    try:
        df = pd.read_excel(src, dtype=str, engine="openpyxl")
    except Exception as e:  # noqa: BLE001
        return [], [f"Could not read the Excel file: {e}"]
    df = df.fillna("")
    df.columns = [str(c).strip().lower() for c in df.columns]
    missing = [c for c in COLUMNS if c.lower() not in df.columns]
    if missing:
        return [], [f"Missing required column(s): {', '.join(missing)}"]
    rows, errors, seen = [], [], set()
    for i, r in df.iterrows():
        line = i + 2
        rec = {k.lower(): str(r[k.lower()]).strip() for k in COLUMNS}
        if not any(rec.values()):
            continue
        problems = []
        if not rec["id"]:
            problems.append("ID is empty")
        elif rec["id"].upper() in seen:
            problems.append(f"duplicate ID {rec['id']}")
        if not rec["name"]:
            problems.append("Name is empty")
        if not _EMAIL.match(rec["email"]):
            problems.append("invalid Email")
        if not rec["role"]:
            problems.append("Role is empty")
        if not rec["skills"]:
            problems.append("Skills is empty")
        if problems:
            errors.append(f"Row {line}: " + "; ".join(problems))
            continue
        rec["id"] = rec["id"].upper()
        seen.add(rec["id"])
        rows.append(rec)
    return rows, errors
